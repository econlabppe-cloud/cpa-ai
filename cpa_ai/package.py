"""חבילת חודש לרואה החשבון: CSV מסכם (נפתח באקסל בעברית), הדוח המוכן, ומסמכי המקור.

המטרה: שהלקוח ילחץ כפתור אחד ויעביר למייצג הכול מסודר, במקום מאות הודעות וקבצים.
"""
from __future__ import annotations

import csv
import io
import json
import re
import zipfile
from datetime import date

from sqlalchemy import select
from sqlalchemy.orm import Session

from . import storage
from .db import Business, Document, Report

COLUMNS = ["תאריך", "כיוון", "ספק / לקוח", "מספר עוסק", "מספר מסמך", "סוג", 'לפני מע"מ', 'מע"מ', 'סה"כ',
           "קטגוריה", "הוכר למס הכנסה", 'מע"מ מנוכה', "סטטוס", "מקור", "קובץ מקור"]
STATUS = {"confirmed": "מאושר", "pending_review": "ממתין לבדיקה", "rejected": "נדחה"}
SOURCE = {"manual": "ידני", "api": "ידני", "upload": "העלאה", "email": "מייל", "whatsapp": "וואטסאפ"}


def _cell(value):
    """מונע CSV injection: שם ספק שמתחיל ב-= + - @ היה מורץ כנוסחה באקסל. מקדימים גרש כדי שיישאר טקסט."""
    if isinstance(value, str) and value[:1] in ("=", "+", "-", "@", "\t", "\r"):
        return "'" + value
    return value


def annual_workbook(s: Session, biz: Business, year: int) -> bytes:
    """חוברת Excel שנתית (RTL): סיכום, הכנסות, הוצאות, מע"מ לפי תקופה, פחת — לעיון המשרד או הלקוח."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font

    from .reports import periodic
    start, end = date(year, 1, 1), date(year, 12, 31)
    est = periodic.build_annual_estimate(s, biz, year, 2.25)
    docs = [d for d in periodic._docs(s, biz.id, start, end) if d.status == "confirmed"]
    wb = Workbook()
    bold = Font(bold=True)

    def sheet(ws, header, rows):
        ws.sheet_view.rightToLeft = True
        ws.append(header)
        for c in ws[1]:
            c.font = bold
            c.alignment = Alignment(horizontal="right")
        for r in rows:
            ws.append([_cell(v) for v in r])
        for col in ws.columns:
            ws.column_dimensions[col[0].column_letter].width = min(max(len(str(c.value or "")) for c in col) + 3, 40)
        ws.freeze_panes = "A2"

    ws = wb.active
    ws.title = "סיכום"
    sheet(ws, ["פריט", "סכום"], [
        ["עסק", biz.name], ["שנה", year],
        ["הכנסות (לפני מע\"מ)", est["revenue_pre_vat"]], ["הוצאות מוכרות", est["recognized_expenses_total"]],
        ["רווח עסקי", est["business_profit"]], ["מס הכנסה מוערך", est["estimated_income_tax"]],
        ["ביטוח לאומי ובריאות מוערך", est["social_security"]["total"]], ["נטל כולל מוערך", est["total_burden"]],
        ["שולם מראש", est["prepaid"]["total"]], ["יתרת מס (שלילי = החזר)", est["income_tax_balance"]],
        *[[f"הוצאות — {k}", v] for k, v in est["recognized_expenses_by_category"].items()],
        ["הערה", "הערכה בלבד — אינה דוח שנתי; פרמטרי המס לא כולם אומתו"]])

    income_rows, expense_rows = [], []
    for d in docs:
        e, dec = d.extracted, d.decision
        if d.direction == "income":
            income_rows.append([d.doc_date.isoformat(), e.get("doc_number"), e.get("customer_name"),
                                e.get("customer_id"), e.get("doc_type"), dec.get("pre_vat_amount"),
                                dec.get("vat_amount"), dec.get("total_amount"), e.get("withholding_tax") or 0])
        else:
            expense_rows.append([d.doc_date.isoformat(), e.get("doc_number"), e.get("vendor_name"),
                                 e.get("vendor_id"), dec.get("category_label"), dec.get("pre_vat_amount"),
                                 dec.get("vat_amount"), dec.get("total_amount"), dec.get("recognized_expense"),
                                 dec.get("deductible_vat")])
    sheet(wb.create_sheet("הכנסות"), ["תאריך", "מספר", "לקוח", "מספר עוסק", "סוג", "לפני מע\"מ", "מע\"מ", "סה\"כ",
                                      "ניכוי במקור"], income_rows)
    sheet(wb.create_sheet("הוצאות"), ["תאריך", "מספר", "ספק", "מספר עוסק", "קטגוריה", "לפני מע\"מ", "מע\"מ", "סה\"כ",
                                      "הוכר", "מע\"מ מנוכה"], expense_rows)
    vat_rows, seen = [], set()
    for m in range(1, 13):
        ps, pe = periodic.vat_period_bounds(year, m, biz.vat_frequency)
        if ps in seen:
            continue
        seen.add(ps)
        p = periodic.build_vat_period(s, biz, ps, pe)
        vat_rows.append([f"{ps.isoformat()} – {pe.isoformat()}", p["sales"]["taxable_pre_vat"], p["sales"]["output_vat"],
                         p["inputs"]["fixed_assets_vat"] + p["inputs"]["other_vat"], p["vat_payable"]])
    sheet(wb.create_sheet("מע\"מ לפי תקופה"), ["תקופה", "עסקאות לפני מע\"מ", "מע\"מ עסקאות", "מע\"מ תשומות", "לתשלום / (להחזר)"],
          vat_rows)
    sheet(wb.create_sheet("פחת"), ["נכס", "נרכש", "קבוצה", "שיעור", "עלות", "פחת השנה"],
          [[x["vendor"], x["purchased"], x["asset_class"], x["rate"], x["cost"], x["depreciation"]]
           for x in est["depreciation"]])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _safe(name: str) -> str:
    return re.sub(r"[^\w.\-֐-׿]+", "_", name)[:80] or "file"


def build_package(s: Session, biz: Business, start: date, end: date) -> bytes:
    docs = list(s.scalars(select(Document).where(
        Document.business_id == biz.id, Document.doc_date >= start, Document.doc_date <= end,
        Document.status != "rejected").order_by(Document.doc_date, Document.id)))
    report = s.scalar(select(Report).where(
        Report.business_id == biz.id, Report.period_start == start, Report.period_end == end,
        Report.status.in_(("prepared", "approved"))).order_by(Report.id.desc()))

    csv_buf = io.StringIO()
    w = csv.writer(csv_buf)
    w.writerow(COLUMNS)
    missing_files: list[int] = []
    zbuf = io.BytesIO()
    with zipfile.ZipFile(zbuf, "w", zipfile.ZIP_DEFLATED) as z:
        for d in docs:
            e, dec = d.extracted, d.decision
            fname = ""
            if d.attachment_path:
                try:
                    fname = f"originals/{d.id:05d}_{_safe(d.attachment_name or 'file')}"
                    z.writestr(fname, storage.read(d.attachment_path))
                except (OSError, ValueError, storage.StorageError):
                    fname, _ = "", missing_files.append(d.id)
            w.writerow([_cell(v) for v in [
                d.doc_date.isoformat() if d.doc_date else "", "הכנסה" if d.direction == "income" else "הוצאה",
                e.get("vendor_name") or e.get("customer_name") or "", e.get("vendor_id") or e.get("customer_id") or "",
                e.get("doc_number") or "", e.get("doc_type") or "", dec.get("pre_vat_amount"), dec.get("vat_amount"),
                dec.get("total_amount"), dec.get("category_label", ""), dec.get("recognized_expense", ""),
                dec.get("deductible_vat", ""), STATUS.get(d.status, d.status), SOURCE.get(d.source, d.source), fname]])
        # utf-8-sig: אקסל פותח עברית נכון
        z.writestr("summary.csv", ("﻿" + csv_buf.getvalue()).encode("utf-8"))
        if report is not None:
            z.writestr("report.json", json.dumps({
                "id": report.id, "kind": report.kind, "status": report.status, "content_hash": report.content_hash,
                "approved_by": report.approved_by, "approver_license": [report.approver_license_type, report.approver_license_number],
                "submitted_at": report.submitted_at.isoformat() if report.submitted_at else None,
                "submission_ref": report.submission_ref, "payload": report.payload}, ensure_ascii=False, indent=2).encode("utf-8"))
        pending = sum(1 for d in docs if d.status == "pending_review")
        lines = [f"חבילת חודש — {biz.name} (עוסק {biz.vat_id})", f"תקופה: {start.isoformat()} עד {end.isoformat()}",
                 f"מסמכים: {len(docs)} · ממתינים לבדיקה: {pending}",
                 "הדוח בחבילה: " + (f"#{report.id} ({report.status})" if report else "לא הוכן"), ""]
        if pending:
            lines.append("שימו לב: יש מסמכים שעוד לא נבדקו ולכן לא נכללים בדוח.")
        if missing_files:
            lines.append(f"קבצי מקור חסרים עבור מסמכים: {missing_files}")
        lines += ["", "הנתונים הוכנו אוטומטית ואינם ייעוץ מס. האחריות המקצועית על ההגשה היא של המייצג."]
        z.writestr("README.txt", "\n".join(lines).encode("utf-8"))
    return zbuf.getvalue()
