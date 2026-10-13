"""ייבוא תנועות בנק ואשראי (CSV / Excel / HTML-בתחפושת-xls / API) והתאמה לחשבוניות.

המטרה: לזהות מה חסר ללקוח. תנועת חיוב בלי קבלה = הוצאה שאולי לא מתועדת; חשבונית בלי תשלום בבנק = משהו לבדוק.
ההתאמה שמרנית: סכום זהה + חלון תאריכים + יחידות. התאמה עמומה לא נעשית אוטומטית.
"""
from __future__ import annotations

import csv
import hashlib
import io
import re
from collections import Counter
from datetime import date, datetime, timedelta
from html.parser import HTMLParser

from sqlalchemy import select
from sqlalchemy.orm import Session

from .db import BankTransaction, Document

MIN_RECEIPT_AMOUNT = 30.0   # תנועות קטנות יותר לא מוצגות כ"חסרה קבלה"
DAYS_BEFORE = 3              # התשלום יכול להקדים מעט את תאריך המסמך
DAYS_AFTER = 60              # אשראי ספקים

# סדר עדיפות: בדף אשראי "תאריך עסקה" קודם ל"תאריך חיוב", ו"סכום חיוב" (בש"ח) קודם ל"סכום עסקה" (במטבע המקור)
HEADERS = {
    "date": ("תאריך עסקה", "תאריך רכישה", "transaction date", "תאריך", "תאריך ערך", "תאריך פעולה", "date",
             "value date"),
    "desc": ("שם בית העסק", "שם בית עסק", "בית עסק", "merchant", "תיאור", "תיאור הפעולה", "פרטים", "הפעולה",
             "description", "details"),
    "debit": ("חובה", "debit", "חיוב"),
    "credit": ("זכות", "credit", "זיכוי"),
    "amount": ("סכום חיוב", "סכום החיוב", "סכום חיוב בש\"ח", "charge amount", "סכום", "amount", "סכום בש\"ח",
               "סכום עסקה", "סכום העסקה"),
}
# כותרת שמזכירה "בית עסק" = דף כרטיס אשראי: הסכומים בו חיוביים לחיוב, ולכן מתהפכים
CARD_MARKERS = ("בית העסק", "בית עסק", "merchant")


class BankCsvError(ValueError):
    pass


def parse_amount(raw: str | None) -> float | None:
    if raw is None:
        return None
    t = raw.strip().replace("₪", "").replace(",", "").replace(" ", "")
    if not t:
        return None
    neg = t.startswith("(") and t.endswith(")") or t.endswith("-")
    t = t.strip("()").rstrip("-")
    try:
        v = float(t)
    except ValueError:
        return None
    return -abs(v) if neg else v


def parse_date(raw: str) -> date | None:
    t = raw.strip()
    for fmt in ("%d/%m/%Y", "%d.%m.%Y", "%Y-%m-%d", "%d/%m/%y", "%d.%m.%y", "%d-%m-%Y"):
        try:
            return datetime.strptime(t, fmt).date()
        except ValueError:
            continue
    return None


def _find(header: list[str], names: tuple[str, ...]) -> int | None:
    """התאמה מדויקת לפי סדר העדיפות, ואחריה התאמת תחילית ("סכום חיוב ₪" ← "סכום חיוב")."""
    low = [str(h or "").strip().lower() for h in header]
    for n in names:
        if n.lower() in low:
            return low.index(n.lower())
    for n in names:
        for i, h in enumerate(low):
            if h.startswith(n.lower() + " "):
                return i
    return None


def parse_csv(text: str) -> list[tuple[date, str, float]]:
    return parse_rows(_csv_rows(text))[0]


def _csv_rows(text: str) -> list[list[str]]:
    text = text.lstrip("\ufeff")
    sample = text[:2000]
    delimiter = ";" if sample.count(";") > sample.count(",") else ("\t" if sample.count("\t") > sample.count(",") else ",")
    return list(csv.reader(io.StringIO(text), delimiter=delimiter))


def _xlsx_rows(raw: bytes) -> list[list[str]]:
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    out = []
    for row in wb.worksheets[0].iter_rows(values_only=True):
        out.append([v.strftime("%d/%m/%Y") if isinstance(v, (datetime, date)) else ("" if v is None else str(v))
                    for v in row])
    return out


class _TableParser(HTMLParser):
    def __init__(self):
        super().__init__()
        self.rows, self._row, self._cell = [], None, None

    def handle_starttag(self, tag, attrs):
        if tag == "tr":
            self._row = []
        elif tag in ("td", "th") and self._row is not None:
            self._cell = []

    def handle_endtag(self, tag):
        if tag in ("td", "th") and self._row is not None and self._cell is not None:
            self._row.append(" ".join("".join(self._cell).split()))
            self._cell = None
        elif tag == "tr" and self._row is not None:
            self.rows.append(self._row)
            self._row = None

    def handle_data(self, data):
        if self._cell is not None:
            self._cell.append(data)


def _decode(raw: bytes) -> str:
    try:
        return raw.decode("utf-8-sig")
    except UnicodeDecodeError:
        return raw.decode("cp1255", errors="replace")  # קידוד נפוץ בייצוא מבנקים ישראליים


def parse_statement(raw: bytes) -> tuple[list[tuple[date, str, float]], str]:
    """דף בנק או אשראי בכל פורמט נפוץ. מחזיר (תנועות, "bank"/"card")."""
    if raw[:4] == b"PK\x03\x04":
        rows = _xlsx_rows(raw)
    elif raw[:8] == b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1":
        raise BankCsvError("קובץ Excel בפורמט ישן (xls) — פתחו אותו ושמרו כ-xlsx או CSV")
    else:
        text = _decode(raw)
        if text.lstrip()[:1] == "<":  # חברות אשראי מייצאות לעיתים טבלת HTML בסיומת xls
            p = _TableParser()
            p.feed(text)
            rows = p.rows
        else:
            rows = _csv_rows(text)
    return parse_rows(rows)


def parse_rows(rows: list[list[str]]) -> tuple[list[tuple[date, str, float]], str]:
    # שורת הכותרת: הראשונה שמכילה עמודת תאריך
    for i, row in enumerate(rows):
        if _find(row, HEADERS["date"]) is not None:
            header, body = row, rows[i + 1:]
            break
    else:
        raise BankCsvError("לא נמצאה שורת כותרת עם עמודת תאריך")
    i_date, i_desc = _find(header, HEADERS["date"]), _find(header, HEADERS["desc"])
    i_debit, i_credit, i_amt = (_find(header, HEADERS[k]) for k in ("debit", "credit", "amount"))
    if i_amt is None and i_debit is None and i_credit is None:
        raise BankCsvError("לא נמצאה עמודת סכום (סכום / חובה / זכות)")
    kind = "card" if any(m in str(h) for h in header for m in CARD_MARKERS) else "bank"

    out = []
    for row in body:
        if len(row) <= i_date or not row[i_date].strip():
            continue
        d = parse_date(row[i_date])
        if d is None:
            continue
        if i_amt is not None:
            amount = parse_amount(row[i_amt] if i_amt < len(row) else None)
        else:
            debit = parse_amount(row[i_debit]) if i_debit is not None and i_debit < len(row) else None
            credit = parse_amount(row[i_credit]) if i_credit is not None and i_credit < len(row) else None
            amount = (credit or 0.0) - abs(debit or 0.0) if (debit is not None or credit is not None) else None
        if not amount:
            continue
        if kind == "card":
            amount = -amount
        desc = row[i_desc].strip() if i_desc is not None and i_desc < len(row) else ""
        out.append((d, desc, round(amount, 2)))
    return out, kind


def import_csv(s: Session, business_id: int, text: str) -> dict:
    rows, kind = parse_rows(_csv_rows(text))
    return import_rows(s, business_id, rows, kind)


def import_rows(s: Session, business_id: int, parsed: list[tuple], account: str = "bank",
                external_ids: list[str | None] | None = None) -> dict:
    """external_ids (מ-API): מזהה יציב מהמקור לדה-דופליקציה. בלעדיו — לפי תאריך, תיאור, סכום ומספר מופע."""
    seen: Counter = Counter()
    added = skipped = 0
    for i, (d, desc, amount) in enumerate(parsed):
        ext = external_ids[i] if external_ids else None
        if ext:
            h = hashlib.sha256(f"ext|{account}|{ext}".encode()).hexdigest()
        else:
            key = f"{d.isoformat()}|{desc}|{amount:.2f}"
            seen[key] += 1  # תנועות זהות באותו קובץ (שתי קפה באותו יום) נשארות נפרדות
            h = hashlib.sha256(f"{key}|{seen[key]}".encode()).hexdigest()
        if s.scalar(select(BankTransaction.id).where(BankTransaction.business_id == business_id,
                                                     BankTransaction.external_hash == h)):
            skipped += 1
            continue
        s.add(BankTransaction(business_id=business_id, tx_date=d, description=desc[:300], amount=amount,
                              account=account, external_hash=h))
        added += 1
    s.flush()
    matched = auto_match(s, business_id)
    s.commit()
    return {"rows": len(parsed), "imported": added, "skipped_duplicates": skipped, "auto_matched": matched,
            "account": account}


def _tokens(text: str) -> set[str]:
    return {t for t in re.split(r"[^\w֐-׿]+", (text or "").lower()) if len(t) >= 2}


def money_flow(d: Document) -> int:
    """כיוון התנועה הצפויה בבנק: הכנסה נכנסת, הוצאה יוצאת, וחשבונית זיכוי הופכת את הכיוון."""
    sign = 1 if d.direction == "income" else -1
    return -sign if d.extracted.get("doc_type") == "credit_invoice" else sign


def auto_match(s: Session, business_id: int) -> int:
    txs = list(s.scalars(select(BankTransaction).where(
        BankTransaction.business_id == business_id, BankTransaction.status == "open")))
    docs = list(s.scalars(select(Document).where(Document.business_id == business_id, Document.status != "rejected",
                                                 Document.doc_date.is_not(None))))
    taken = {t.matched_document_id for t in s.scalars(select(BankTransaction).where(
        BankTransaction.business_id == business_id, BankTransaction.matched_document_id.is_not(None)))}
    matched = 0
    for tx in sorted(txs, key=lambda t: t.tx_date):
        cands = []
        for d in docs:
            total = d.decision.get("total_amount")
            if d.id in taken or total is None or money_flow(d) * tx.amount <= 0:
                continue
            if abs(abs(total) - abs(tx.amount)) > 0.01:
                continue
            delta = (tx.tx_date - d.doc_date).days
            if not (-DAYS_BEFORE <= delta <= DAYS_AFTER):
                continue
            name = d.extracted.get("vendor_name") or d.extracted.get("customer_name") or ""
            overlap = len(_tokens(name) & _tokens(tx.description))
            cands.append((overlap, -abs(delta), d))
        if not cands:
            continue
        cands.sort(key=lambda c: (c[0], c[1]), reverse=True)
        if len(cands) > 1 and cands[0][:2] == cands[1][:2]:
            continue  # עמום: לא מנחשים
        tx.matched_document_id, tx.status = cands[0][2].id, "matched"
        taken.add(cands[0][2].id)
        matched += 1
    s.flush()
    return matched


def unmatched(s: Session, business_id: int) -> dict:
    txs = list(s.scalars(select(BankTransaction).where(BankTransaction.business_id == business_id)
                         .order_by(BankTransaction.tx_date.desc())))
    open_debits = [t for t in txs if t.status == "open" and t.amount < 0 and abs(t.amount) >= MIN_RECEIPT_AMOUNT]
    no_payment: list[dict] = []
    if txs:
        lo, hi = min(t.tx_date for t in txs), max(t.tx_date for t in txs)
        paid = {t.matched_document_id for t in txs if t.matched_document_id}
        for d in s.scalars(select(Document).where(
                Document.business_id == business_id, Document.direction == "expense", Document.status == "confirmed",
                Document.doc_date.is_not(None))):
            # רק מסמכים שהבנק מכסה את תאריכם (בתוספת חלון אשראי), כדי לא להציף כשיש רק חודש אחד בנק
            if d.id not in paid and lo <= d.doc_date <= hi - timedelta(days=DAYS_AFTER // 2):
                no_payment.append({"document_id": d.id, "vendor": d.extracted.get("vendor_name"),
                                   "date": d.doc_date.isoformat(), "total": d.decision.get("total_amount")})
    return {
        "transactions_without_document": [
            {"id": t.id, "date": t.tx_date.isoformat(), "description": t.description, "amount": t.amount,
             "account": t.account or "bank"}
            for t in open_debits],
        "documents_without_payment": no_payment,
    }
